from calendar import monthrange
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db.models import Max, Min, Sum
from django.shortcuts import redirect, render
from django.http import HttpResponse, JsonResponse

from .feriados import (
    FERIADOS_2026,
    TIPOS_RESTRINGIDOS,
    dia_habil_para,
    siguiente_dia_habil,
    mensaje_no_habil,
)
from .models import (
    ROLES,
    TIPOS_LOCAL,
    Historial,
    Local,
    LoginLog,
    Pago,
    PagoDiario,
    Perfil,
    ValorUTM,
)

MESES_ES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

DIAS_SEMANA_ES = [
    "Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo",
]

DIAS_GENERACION = 5
PREFIJO_FOLIO_GEN = "GEN"
UTM_BASE = Decimal("66000")


def _mes_num(valor):
    if not valor:
        return None
    if valor.isdigit():
        return int(valor)
    buscado = valor.strip().lower()
    for indice, nombre_mes in enumerate(MESES_ES):
        if nombre_mes.lower() == buscado:
            return indice + 1
    return None


def registrar_historial(request, accion, tipo, folio, local_numero, monto):
    Historial.objects.create(
        usuario=request.user,
        accion=accion,
        tipo=tipo,
        folio=folio,
        local_numero=local_numero,
        monto=monto,
    )


def login_view(request):
    if request.user.is_authenticated:
        return redirect("inicio")
    if request.method == "POST":
        username = request.POST.get("username", "")
        password = request.POST.get("password", "")
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            return redirect("inicio")
        messages.error(request, "Usuario o contraseña incorrectos.")
    return render(request, "core/login.html")


def logout_view(request):
    logout(request)
    return redirect("login")


@login_required
def inicio(request):
    if request.user.perfil.rol == "administrador":
        return redirect("usuarios_admin")
    return redirect("pagos")


@login_required
def pagos(request, tipo="local"):
    if request.user.perfil.rol == "administrador":
        messages.error(request, "Tu perfil no tiene acceso al módulo de pagos.")
        return redirect("usuarios_admin")

    tipos_validos = {v for v, _ in TIPOS_LOCAL}
    if tipo not in tipos_validos:
        tipo = "local"

    numero_filtro = request.GET.get("numero", "").strip()
    mes_desde = request.GET.get("mes_desde", "").strip()
    mes_hasta = request.GET.get("mes_hasta", "").strip()
    anio_desde = request.GET.get("anio_desde", "").strip()
    anio_hasta = request.GET.get("anio_hasta", "").strip()
    folio_filtro = request.GET.get("folio", "").strip()
    empresario_filtro = request.GET.get("empresario", "").strip()

    pagos_lista = Pago.objects.filter(local__tipo=tipo)
    if numero_filtro:
        pagos_lista = pagos_lista.filter(local__numero__icontains=numero_filtro)
    if folio_filtro:
        pagos_lista = pagos_lista.filter(folio__icontains=folio_filtro)
    if empresario_filtro:
        pagos_lista = pagos_lista.filter(empresario__icontains=empresario_filtro)

    mes_desde_num = _mes_num(mes_desde)
    mes_hasta_num = _mes_num(mes_hasta)
    if mes_desde_num:
        pagos_lista = pagos_lista.filter(fecha__month__gte=mes_desde_num)
    if mes_hasta_num:
        pagos_lista = pagos_lista.filter(fecha__month__lte=mes_hasta_num)
    if anio_desde.isdigit():
        pagos_lista = pagos_lista.filter(fecha__year__gte=int(anio_desde))
    if anio_hasta.isdigit():
        pagos_lista = pagos_lista.filter(fecha__year__lte=int(anio_hasta))
    pagos_lista = pagos_lista.select_related("local").order_by("-fecha")

    años_disponibles = (
        Pago.objects.filter(local__tipo=tipo)
        .dates("fecha", "year")
    )
    años_lista = sorted({a.year for a in años_disponibles})
    meses_lista = MESES_ES

    base_tipo = Pago.objects.filter(local__tipo=tipo)
    numeros_sugeridos = list(
        base_tipo.values_list("local__numero", flat=True).distinct().order_by("local__numero")
    )
    folios_sugeridos = list(
        base_tipo.values_list("folio", flat=True).distinct().order_by("folio")
    )
    empresarios_sugeridos = list(
        base_tipo.values_list("empresario", flat=True).distinct().order_by("empresario")
    )

    es_caja = request.user.perfil.rol == "caja"
    es_admin = request.user.perfil.rol == "administrador"

    # Precalcular desglose por día: solo los días que pertenecen al mes del pago
    for p in pagos_lista:
        detalles = list(p.detalles_diarios.all())
        por_fecha = {d.fecha: d for d in detalles}
        ultimo_dia = monthrange(p.fecha.year, p.fecha.month)[1]
        desglose = []
        for dia in range(1, ultimo_dia + 1):
            dia_fecha = date(p.fecha.year, p.fecha.month, dia)
            if not dia_habil_para(tipo, dia_fecha):
                continue
            d = por_fecha.get(dia_fecha)
            desglose.append({
                "id": d.id if d else None,
                "dia_nombre": DIAS_SEMANA_ES[dia_fecha.weekday()],
                "dia_num": dia_fecha.day,
                "fecha": dia_fecha.isoformat(),
                "folio": (d.folio or "") if d else "",
                "valor": d.valor if d else Decimal("0"),
            })
        p.desglose = desglose
        p.desglose_total = sum((d.valor for d in detalles), Decimal("0"))

    # Valor UTM visible (para el indicador/card): mes del filtro o mes actual
    anio_actual = date.today().year
    mes_visible = mes_desde_num or date.today().month

    # Empresarios nuevos del mes visible: primer pago en ese mes (por tipo)
    empresarios_nuevos = (
        Pago.objects.filter(local__tipo=tipo)
        .exclude(empresario="")
        .values("empresario")
        .annotate(primera=Min("fecha"))
        .filter(primera__year=anio_actual, primera__month=mes_visible)
        .count()
    )

    context = {
        "tipos": TIPOS_LOCAL,
        "tipo_actual": tipo,
        "pagos": pagos_lista,
        "numero_filtro": numero_filtro,
        "mes_desde": mes_desde,
        "mes_hasta": mes_hasta,
        "anio_desde": anio_desde,
        "anio_hasta": anio_hasta,
        "folio_filtro": folio_filtro,
        "empresario_filtro": empresario_filtro,
        "años_lista": años_lista,
        "meses_lista": meses_lista,
        "numeros_sugeridos": numeros_sugeridos,
        "folios_sugeridos": folios_sugeridos,
        "empresarios_sugeridos": empresarios_sugeridos,
        "es_caja": es_caja,
        "es_admin": es_admin,
        "generacion_disponible": es_caja and date.today().day <= DIAS_GENERACION,
        "empresarios_nuevos": empresarios_nuevos,
        "mes_nuevos": MESES_ES[mes_visible - 1],
        "anio_nuevos": anio_actual,
    }
    return render(request, "core/pagos.html", context)


def _datos_generacion(hoy):
    periodo = f"{hoy.year}{hoy.month:02d}"
    filas = []
    for local in Local.objects.filter(ocupado=True).order_by("numero"):
        ultimo = local.pagos.order_by("-fecha", "-id").first()
        referencia = local.pagos.aggregate(v=Max("valor"))["v"] or Decimal("0")
        filas.append({
            "local": local,
            "ultimo": ultimo,
            "referencia": referencia,
            "folio": f"{PREFIJO_FOLIO_GEN}-{periodo}-{local.numero}",
            "existe": Pago.objects.filter(
                folio=f"{PREFIJO_FOLIO_GEN}-{periodo}-{local.numero}"
            ).exists(),
            "fecha_aplicada": siguiente_dia_habil(local.tipo, hoy),
        })
    return filas


@login_required
def generar_deuda(request):
    if request.user.perfil.rol != "caja":
        messages.error(request, "No tienes permisos para generar la deuda.")
        return redirect("pagos")

    hoy = date.today()
    if hoy.day > DIAS_GENERACION:
        messages.error(
            request,
            "La generación automática solo está disponible los primeros 5 días del mes.",
        )
        return redirect("pagos")

    filas = _datos_generacion(hoy)

    if request.method == "POST":
        creados = 0
        omitidos = 0
        total = Decimal("0")
        for fila in filas:
            if fila["existe"]:
                continue
            valor_post = request.POST.get(f"valor_{fila['local'].id}", "").strip()
            try:
                valor_dec = Decimal(valor_post) if valor_post else None
            except InvalidOperation:
                valor_dec = None
            if valor_dec is None or valor_dec <= 0:
                omitidos += 1
                continue
            if Pago.objects.filter(folio=fila["folio"]).exists():
                omitidos += 1
                continue
            local = fila["local"]
            ultimo = fila["ultimo"]
            fecha_obj = fila["fecha_aplicada"]
            Pago.objects.create(
                local=local,
                fecha=fecha_obj,
                mes=MESES_ES[fecha_obj.month - 1],
                empresario=ultimo.empresario if ultimo else "",
                rut=ultimo.rut if ultimo else "",
                descripcion=f"Deuda mensual {hoy.year}-{hoy.month:02d} generada automáticamente",
                valor=valor_dec,
                folio=fila["folio"],
            )
            registrar_historial(
                request, "generar", local.tipo, fila["folio"], local.numero, valor_dec
            )
            creados += 1
            total += valor_dec
        if creados:
            messages.success(
                request,
                f"Deuda del mes generada: {creados} registro(s) por ${total:,.0f}."
                f" Omitidos: {omitidos}.",
            )
        else:
            messages.error(request, "No se generó ningún registro nuevo.")
        return redirect("pagos_tipo", tipo="local")

    pendientes = [f for f in filas if not f["existe"]]
    ya_generados = [f for f in filas if f["existe"]]
    total_a_generar = sum((f["referencia"] for f in pendientes), Decimal("0"))
    context = {
        "filas": filas,
        "pendientes": pendientes,
        "ya_generados": ya_generados,
        "total_a_generar": total_a_generar,
        "periodo": f"{hoy.year}-{hoy.month:02d}",
    }
    return render(request, "core/generar_deuda.html", context)


@login_required
def importar_excel_pagos(request):
    if request.user.perfil.rol != "caja":
        messages.error(request, "No tienes permisos para importar pagos.")
        return redirect("pagos")

    if request.method == "POST":
        archivo = request.FILES.get("archivo")
        if not archivo:
            messages.error(request, "Selecciona un archivo Excel (.xlsx).")
            return redirect("importar_excel_pagos")

        from .importador import importar_excel_desde_bytes

        r = importar_excel_desde_bytes(archivo, request.POST.get("tipo") or None)
        if r.get("error"):
            messages.error(request, r["error"])
        else:
            for msg in r.get("fila", [])[:8]:
                messages.warning(request, msg)
            messages.success(
                request,
                f"Importación terminada: {r['creados']} creados, {r['actualizados']} actualizados, "
                f"{r['sin_cambios']} sin cambios, {r['errores']} errores.",
            )
            if r["errores"]:
                messages.error(request, f"Hubo {r['errores']} errores; revisa los avisos.")
        return redirect("pagos")

    context = {
        "tipos": TIPOS_LOCAL,
    }
    return render(request, "core/importar_excel.html", context)


@login_required
def agregar_pago(request, tipo):
    if request.user.perfil.rol != "caja":
        messages.error(request, "No tienes permisos para registrar pagos.")
        return redirect("pagos_tipo", tipo=tipo)

    if request.method == "POST":
        try:
            numero = request.POST.get("numero", "").strip()
            fecha = request.POST.get("fecha", "")
            mes = request.POST.get("mes", "").strip()
            empresario = request.POST.get("empresario", "").strip()
            rut = request.POST.get("rut", "").strip()
            descripcion = request.POST.get("descripcion", "").strip()
            tipo_deuda = request.POST.get("tipo_deuda", "").strip()
            domicilio = request.POST.get("domicilio", "").strip()
            fono = request.POST.get("fono", "").strip()
            correo = request.POST.get("correo", "").strip()
            periodo = request.POST.get("periodo", "").strip()
            valor = request.POST.get("valor", "")
            fecha_pago = request.POST.get("fecha_pago", "") or None
            folio = request.POST.get("folio", "").strip()

            if not numero or not fecha:
                raise ValueError("N° de local y fecha son obligatorios.")
            if not valor:
                raise ValueError("El valor es obligatorio.")
            if not folio:
                raise ValueError("El folio es obligatorio.")

            from datetime import datetime
            fecha_obj = datetime.strptime(fecha, "%Y-%m-%d").date()

            if not dia_habil_para(tipo, fecha_obj):
                raise ValueError(mensaje_no_habil(tipo, fecha_obj))

            if not mes:
                mes = MESES_ES[fecha_obj.month - 1]

            valor_dec = Decimal(valor)
            if valor_dec < 0:
                raise ValueError("El valor no puede ser negativo.")

            local, _ = Local.objects.get_or_create(
                numero=numero, defaults={"tipo": tipo, "ocupado": True}
            )
            if local.tipo != tipo:
                local.tipo = tipo
                local.save()

            if Pago.objects.filter(folio=folio).exists():
                raise ValueError(f"Ya existe un registro con el folio {folio}.")

            fecha_pago_obj = None
            if fecha_pago:
                fecha_pago_obj = datetime.strptime(fecha_pago, "%Y-%m-%d").date()

            pago = Pago.objects.create(
                local=local,
                fecha=fecha_obj,
                mes=mes,
                empresario=empresario,
                rut=rut,
                descripcion=descripcion,
                tipo_deuda=tipo_deuda,
                domicilio=domicilio,
                fono=fono,
                correo=correo,
                periodo=periodo,
                valor=valor_dec,
                fecha_pago=fecha_pago_obj,
                folio=folio,
            )
            repartir_valor_dias(pago, tipo, fecha_obj, valor_dec)
            registrar_historial(
                request, "crear", tipo, folio, numero, valor_dec
            )
            messages.success(request, "Registro guardado correctamente.")
            return redirect("pagos_tipo", tipo=tipo)
        except ValueError as e:
            messages.error(request, str(e))
        except Exception as e:
            messages.error(request, f"No se pudo guardar: {e}")

    context = {
        "tipo": tipo,
        "hoy": date.today().isoformat(),
        "restringido": tipo in TIPOS_RESTRINGIDOS,
        "feriados": [f.isoformat() for f in FERIADOS_2026],
    }
    return render(request, "core/agregar_pago.html", context)


@login_required
def abonar_pago(request, tipo, folio):
    if request.user.perfil.rol != "caja":
        return JsonResponse({"ok": False, "error": "Sin permisos"}, status=403)

    try:
        pago = Pago.objects.select_related("local").get(folio=folio)
    except Pago.DoesNotExist:
        return JsonResponse({"ok": False, "error": "Pago no encontrado"}, status=404)

    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "Método no permitido"}, status=405)

    accion = request.POST.get("accion", "guardar")

    def _fecha_dia():
        valor_fecha = request.POST.get("dia_fecha", "").strip()
        if valor_fecha:
            return datetime.strptime(valor_fecha, "%Y-%m-%d").date()
        valor_dia = int(request.POST.get("dia_num") or 0)
        return date(pago.fecha.year, pago.fecha.month, valor_dia)

    try:
        if accion == "eliminar":
            fecha_dia = _fecha_dia()
            PagoDiario.objects.filter(pago=pago, fecha=fecha_dia).delete()
            return JsonResponse({"ok": True, "total": _total_pago(pago)})

        # accion == guardar
        fecha_dia = _fecha_dia()
        dia_nombre = request.POST.get("dia_nombre", "").strip()
        valor_str = request.POST.get("valor", "").strip()
        folio_dia = request.POST.get("folio", "").strip()

        if not (date(1900, 1, 1) <= fecha_dia <= date(2100, 12, 31)):
            return JsonResponse({"ok": False, "error": "Fecha fuera de rango"}, status=400)

        if not dia_nombre:
            dia_nombre = DIAS_SEMANA_ES[fecha_dia.weekday()]

        if folio_dia:
            en_uso = (
                PagoDiario.objects.filter(folio=folio_dia)
                .exclude(pago=pago, fecha=fecha_dia)
                .exists()
            )
            if en_uso:
                return JsonResponse({"ok": False, "error": "Ese folio ya está en uso."}, status=400)

        valor_dec = Decimal(valor_str) if valor_str else Decimal("0")
        if valor_dec < 0:
            return JsonResponse({"ok": False, "error": "El valor no puede ser negativo"}, status=400)

        PagoDiario.objects.update_or_create(
            pago=pago,
            fecha=fecha_dia,
            defaults={"dia": dia_nombre, "valor": valor_dec, "folio": folio_dia or None},
        )
        return JsonResponse({"ok": True, "total": _total_pago(pago)})
    except (ValueError, InvalidOperation):
        return JsonResponse({"ok": False, "error": "Datos inválidos"}, status=400)


def _total_pago(pago):
    total = sum(
        (d.valor for d in pago.detalles_diarios.all()),
        Decimal("0"),
    )
    return str(total)


def repartir_valor_dias(pago, tipo, fecha, valor):
    """Divide el monto del pago entre los días hábiles del mes y lo guarda
    en cada PagoDiario del desglose. Conserva los días existentes por fecha y
    elimina los que dejan de pertenecer al mes."""
    ultimo_dia = monthrange(fecha.year, fecha.month)[1]
    fechas = []
    for dia_num in range(1, ultimo_dia + 1):
        dia_fecha = date(fecha.year, fecha.month, dia_num)
        if not dia_habil_para(tipo, dia_fecha):
            continue
        fechas.append(dia_fecha)

    n = len(fechas)
    if n == 0:
        pago.detalles_diarios.all().delete()
        return

    centavos_total = int((valor * Decimal("100")).quantize(Decimal("1")))
    base, resto = divmod(centavos_total, n)
    existentes = {d.fecha: d for d in pago.detalles_diarios.all()}
    for i, dia_fecha in enumerate(fechas):
        centavos = base + (1 if i < resto else 0)
        valor_dia = Decimal(centavos) / 100
        d = existentes.pop(dia_fecha, None)
        if d is None:
            PagoDiario.objects.create(
                pago=pago,
                fecha=dia_fecha,
                dia=DIAS_SEMANA_ES[dia_fecha.weekday()],
                valor=valor_dia,
                folio=None,
            )
        else:
            d.dia = DIAS_SEMANA_ES[dia_fecha.weekday()]
            d.valor = valor_dia
            d.save()

    for d in existentes.values():
        d.delete()


@login_required
def editar_pago(request, tipo, folio):
    if request.user.perfil.rol != "caja":
        messages.error(request, "No tienes permisos para editar pagos.")
        return redirect("pagos_tipo", tipo=tipo)

    try:
        pago = Pago.objects.select_related("local").get(folio=folio)
    except Pago.DoesNotExist:
        messages.error(request, "No se encontró el registro con ese folio.")
        return redirect("pagos_tipo", tipo=tipo)

    if request.method == "POST":
        formato = request.POST.get("formato", "")
        try:
            fecha = request.POST.get("fecha", "").strip()
            mes = request.POST.get("mes", "").strip()
            empresario = request.POST.get("empresario", "").strip()
            rut = request.POST.get("rut", "").strip()
            descripcion = request.POST.get("descripcion", "").strip()
            tipo_deuda = request.POST.get("tipo_deuda", "").strip()
            domicilio = request.POST.get("domicilio", "").strip()
            fono = request.POST.get("fono", "").strip()
            correo = request.POST.get("correo", "").strip()
            periodo = request.POST.get("periodo", "").strip()
            valor = request.POST.get("valor", "").strip()
            fecha_pago = request.POST.get("fecha_pago", "").strip() or None
            numero = request.POST.get("numero", "").strip() or pago.local.numero
            folio_nuevo = request.POST.get("folio", "").strip() or pago.folio

            from datetime import datetime

            if not fecha:
                raise ValueError("La fecha es obligatoria.")
            fecha_obj = datetime.strptime(fecha, "%Y-%m-%d").date()

            if not dia_habil_para(tipo, fecha_obj):
                raise ValueError(mensaje_no_habil(tipo, fecha_obj))

            if not mes:
                mes = MESES_ES[fecha_obj.month - 1]

            valor_dec = Decimal(valor)
            if valor_dec < 0:
                raise ValueError("El valor no puede ser negativo.")

            fecha_pago_obj = None
            if fecha_pago:
                fecha_pago_obj = datetime.strptime(fecha_pago, "%Y-%m-%d").date()

            if not numero:
                raise ValueError("El N° de local es obligatorio.")

            if folio_nuevo != pago.folio and Pago.objects.filter(folio=folio_nuevo).exists():
                raise ValueError(f"Ya existe un registro con el folio {folio_nuevo}.")

            fecha_anterior = pago.fecha
            valor_anterior = pago.valor
            local_obj = pago.local
            if numero != pago.local.numero:
                otro = Local.objects.filter(numero=numero).first()
                if otro and otro != pago.local:
                    raise ValueError(f"Ya existe un local con el N° {numero}.")
                local_obj, creado = Local.objects.get_or_create(
                    numero=numero, defaults={"tipo": tipo, "ocupado": True}
                )
                if local_obj.tipo != tipo:
                    local_obj.tipo = tipo
                    local_obj.save()

            viejo_local = pago.local
            pago.local = local_obj
            pago.folio = folio_nuevo
            pago.fecha = fecha_obj
            pago.mes = mes
            pago.empresario = empresario
            pago.rut = rut
            pago.descripcion = descripcion
            pago.tipo_deuda = tipo_deuda
            pago.domicilio = domicilio
            pago.fono = fono
            pago.correo = correo
            pago.periodo = periodo
            pago.valor = valor_dec
            pago.fecha_pago = fecha_pago_obj
            pago.save()

            if (fecha_anterior.year, fecha_anterior.month) != (fecha_obj.year, fecha_obj.month) or \
               valor_anterior != valor_dec:
                repartir_valor_dias(pago, tipo, fecha_obj, valor_dec)
            if viejo_local != local_obj and not viejo_local.pagos.exists():
                viejo_local.ocupado = False
                viejo_local.save()

            registrar_historial(
                request, "editar", pago.local.tipo, pago.folio, pago.local.numero, valor_dec
            )
            if formato == "json":
                return JsonResponse({
                    "ok": True,
                    "numero": pago.local.numero,
                    "fecha": pago.fecha.strftime("%d/%m/%Y"),
                    "fecha_raw": pago.fecha.isoformat(),
                    "mes": pago.mes,
                    "empresario": pago.empresario,
                    "rut": pago.rut,
                    "descripcion": pago.descripcion,
                    "tipo_deuda": pago.tipo_deuda,
                    "domicilio": pago.domicilio,
                    "fono": pago.fono,
                    "correo": pago.correo,
                    "periodo": pago.periodo,
                    "valor": f"{pago.valor:,.0f}",
                    "valor_raw": f"{pago.valor}",
                    "fecha_pago": pago.fecha_pago.strftime("%d/%m/%Y") if pago.fecha_pago else "",
                    "fecha_pago_raw": pago.fecha_pago.isoformat() if pago.fecha_pago else "",
                    "folio": pago.folio,
                })
            messages.success(request, "Registro actualizado correctamente.")
            return redirect("pagos_tipo", tipo=tipo)
        except (ValueError, InvalidOperation) as e:
            if formato == "json":
                return JsonResponse({"ok": False, "error": str(e)}, status=400)
            messages.error(request, f"No se pudo actualizar: {e}")
        except Exception as e:
            if formato == "json":
                return JsonResponse({"ok": False, "error": f"No se pudo guardar: {e}"}, status=400)
            messages.error(request, f"No se pudo actualizar: {e}")

    context = {
        "pago": pago,
        "tipo": tipo,
        "restringido": tipo in TIPOS_RESTRINGIDOS,
        "feriados": [f.isoformat() for f in FERIADOS_2026],
    }
    return render(request, "core/editar_pago.html", context)


@login_required
def eliminar_pago(request, tipo, folio):
    if request.user.perfil.rol != "caja":
        messages.error(request, "No tienes permisos para eliminar pagos.")
        return redirect("pagos_tipo", tipo=tipo)

    formato = request.POST.get("formato", "")
    try:
        pago = Pago.objects.select_related("local").get(folio=folio)
    except Pago.DoesNotExist:
        if formato == "json":
            return JsonResponse({"ok": False, "error": "No se encontró el registro."}, status=404)
        messages.error(request, "No se encontró el registro.")
        return redirect("pagos_tipo", tipo=tipo)

    if request.method == "POST":
        registrar_historial(
            request, "eliminar", pago.local.tipo, pago.folio, pago.local.numero, pago.valor
        )
        pago.delete()
        if formato == "json":
            return JsonResponse({"ok": True})
        messages.success(request, "Registro eliminado correctamente.")
        return redirect("pagos_tipo", tipo=tipo)

    context = {"pago": pago, "tipo": tipo}
    return render(request, "core/eliminar_pago.html", context)


def _registros_historial():
    registros = []
    for h in Historial.objects.select_related("usuario").order_by("-fecha_hora"):
        registros.append(
            {
                "fecha_hora": h.fecha_hora,
                "usuario": h.usuario.username,
                "accion": h.get_accion_display(),
                "local": h.local_numero,
                "folio": h.folio,
                "monto": h.monto,
                "tipo": "operacion",
            }
        )
    for l in LoginLog.objects.select_related("usuario").order_by("-fecha_hora"):
        registros.append(
            {
                "fecha_hora": l.fecha_hora,
                "usuario": l.usuario.username,
                "accion": l.get_accion_display(),
                "local": "—",
                "folio": "—",
                "monto": None,
                "tipo": "ingreso" if l.accion == "login" else "salida",
            }
        )
    registros.sort(key=lambda r: r["fecha_hora"], reverse=True)
    return registros


@login_required
def historial(request):
    if request.user.perfil.rol == "administrador":
        return redirect("usuarios_admin")
    context = {"registros": _registros_historial()}
    return render(request, "core/historial.html", context)


@login_required
def locales_desocupados(request):
    if request.user.perfil.rol == "administrador":
        messages.error(request, "Tu perfil no tiene acceso a este módulo.")
        return redirect("usuarios_admin")
    desocupados = Local.objects.filter(ocupado=False).order_by("numero")
    context = {"desocupados": desocupados}
    return render(request, "core/desocupados.html", context)


@login_required
def usuarios_admin(request):
    if request.user.perfil.rol != "administrador":
        messages.error(request, "No tienes permisos para administrar usuarios.")
        return redirect("inicio")

    if request.method == "POST":
        accion = request.POST.get("accion", "")
        if accion == "crear":
            username = request.POST.get("username", "").strip()
            password = request.POST.get("password", "")
            rol = request.POST.get("rol", "caja")
            try:
                if not username or not password:
                    raise ValueError("Usuario y contraseña son obligatorios.")
                if User.objects.filter(username=username).exists():
                    raise ValueError(f"El usuario '{username}' ya existe.")
                user = User.objects.create_user(username=username, password=password)
                user.perfil.rol = rol
                user.perfil.save()
                messages.success(request, "Usuario creado correctamente.")
            except ValueError as e:
                messages.error(request, str(e))
        elif accion == "cambiar_rol":
            username = request.POST.get("username", "").strip()
            nuevo_rol = request.POST.get("rol", "caja")
            try:
                user = User.objects.get(username=username)
                user.perfil.rol = nuevo_rol
                user.perfil.save()
                messages.success(request, f"Rol de '{username}' actualizado.")
            except User.DoesNotExist:
                messages.error(request, "Usuario no encontrado.")

    usuarios = Perfil.objects.select_related("user").order_by("user__username")
    context = {
        "usuarios": usuarios,
        "roles": ROLES,
        "registros": _registros_historial()[:50],
    }
    return render(request, "core/usuarios.html", context)


@login_required
def reportes(request):
    tipo = request.GET.get("tipo", "")
    anio_desde = request.GET.get("anio_desde", "").strip()
    anio_hasta = request.GET.get("anio_hasta", "").strip()
    mes_desde = request.GET.get("mes_desde", "").strip()
    mes_hasta = request.GET.get("mes_hasta", "").strip()
    exportar = request.GET.get("exportar", "")

    base = Pago.objects.all()
    if tipo:
        base = base.filter(local__tipo=tipo)
    if anio_desde.isdigit():
        base = base.filter(fecha__year__gte=int(anio_desde))
    if anio_hasta.isdigit():
        base = base.filter(fecha__year__lte=int(anio_hasta))
    mes_desde_num = _mes_num(mes_desde)
    mes_hasta_num = _mes_num(mes_hasta)
    if mes_desde_num:
        base = base.filter(fecha__month__gte=mes_desde_num)
    if mes_hasta_num:
        base = base.filter(fecha__month__lte=mes_hasta_num)

    total_general = base.aggregate(total=Sum("valor"))["total"] or Decimal("0")
    total_pendiente = base.filter(valor__gt=0).aggregate(total=Sum("valor"))["total"] or Decimal("0")
    cantidad = base.count()

    resumen_tipos = []
    for valor, nombre in TIPOS_LOCAL:
        qs = base.filter(local__tipo=valor)
        tot = qs.aggregate(t=Sum("valor"))["t"] or Decimal("0")
        cant = qs.count()
        if cant:
            resumen_tipos.append({
                "codigo": valor,
                "nombre": nombre,
                "total": tot,
                "cantidad": cant,
            })

    años_disponibles = [a.year for a in Pago.objects.dates("fecha", "year")]

    context = {
        "tipos": TIPOS_LOCAL,
        "tipo_filtro": tipo,
        "anio_desde": anio_desde,
        "anio_hasta": anio_hasta,
        "mes_desde": mes_desde,
        "mes_hasta": mes_hasta,
        "años_lista": sorted(años_disponibles),
        "meses_lista": MESES_ES,
        "resumen_tipos": resumen_tipos,
        "total_general": total_general,
        "total_pendiente": total_pendiente,
        "cantidad": cantidad,
    }

    if exportar:
        from django.http import HttpResponse
        from openpyxl import Workbook
        from openpyxl.styles import Font

        pagos = list(
            base.select_related("local")
            .prefetch_related("detalles_diarios")
            .order_by("fecha")
        )

        wb = Workbook()

        # ---- Hoja 1: cuentas del mes por empresario ----
        ws_cuentas = wb.active
        ws_cuentas.title = "Cuentas del Mes"
        cabeceras_cuentas = [
            "Mes", "Año", "Empresario", "RUT", "Folio", "N° Local", "Tipo local",
            "Tipo deuda", "Descripción", "Domicilio", "Teléfono", "Correo",
            "Período", "Fecha pago", "Valor", "Estado",
        ]
        ws_cuentas.append(cabeceras_cuentas)
        for celda in ws_cuentas[1]:
            celda.font = Font(bold=True)

        for pago in pagos:
            estado = "Pagado" if pago.valor == 0 else "Pendiente"
            ws_cuentas.append([
                pago.mes,
                pago.fecha.year,
                pago.empresario,
                pago.rut,
                pago.folio,
                pago.local.numero,
                pago.local.get_tipo_display(),
                pago.tipo_deuda,
                pago.descripcion,
                pago.domicilio,
                pago.fono,
                pago.correo,
                pago.periodo,
                pago.fecha_pago.isoformat() if pago.fecha_pago else "",
                float(pago.valor),
                estado,
            ])

        fila_total = [""] * len(cabeceras_cuentas)
        fila_total[0] = "TOTAL GENERAL"
        fila_total[14] = float(total_general)
        ws_cuentas.append([])
        ws_cuentas.append(fila_total)
        ws_cuentas.cell(row=ws_cuentas.max_row, column=1).font = Font(bold=True)
        fila_pend = [""] * len(cabeceras_cuentas)
        fila_pend[0] = "TOTAL PENDIENTE"
        fila_pend[14] = float(total_pendiente)
        ws_cuentas.append(fila_pend)
        ws_cuentas.cell(row=ws_cuentas.max_row, column=1).font = Font(bold=True)
        ws_cuentas.append(["Generado", date.today().isoformat()])

        for col, ancho in zip(
            "ABCDEFGHIJKLMNOP",
            (12, 6, 28, 15, 14, 10, 14, 14, 30, 24, 14, 20, 12, 14, 14, 12),
        ):
            ws_cuentas.column_dimensions[col].width = ancho

        # ---- Hoja 2: desglose diario de todos los pagos ----
        ws_desglose = wb.create_sheet("Desglose Diario")
        cabeceras_desglose = [
            "Folio", "Mes", "Año", "Empresario", "RUT", "N° Local", "Tipo local",
            "Tipo deuda", "Descripción", "Domicilio", "Teléfono", "Correo",
            "Período", "Fecha pago", "Valor total", "Día", "Día de la semana",
            "Fecha día", "Valor diario", "Estado",
        ]
        ws_desglose.append(cabeceras_desglose)
        for celda in ws_desglose[1]:
            celda.font = Font(bold=True)

        for pago in pagos:
            estado = "Pagado" if pago.valor == 0 else "Pendiente"
            for det in pago.detalles_diarios.all():
                ws_desglose.append([
                    pago.folio,
                    pago.mes,
                    pago.fecha.year,
                    pago.empresario,
                    pago.rut,
                    pago.local.numero,
                    pago.local.get_tipo_display(),
                    pago.tipo_deuda,
                    pago.descripcion,
                    pago.domicilio,
                    pago.fono,
                    pago.correo,
                    pago.periodo,
                    pago.fecha_pago.isoformat() if pago.fecha_pago else "",
                    float(pago.valor),
                    det.fecha.day,
                    det.dia,
                    det.fecha.isoformat(),
                    float(det.valor),
                    estado,
                ])

        fila_total2 = [""] * len(cabeceras_desglose)
        fila_total2[0] = "TOTAL GENERAL"
        fila_total2[14] = float(total_general)
        fila_total2[18] = float(total_general)
        ws_desglose.append([])
        ws_desglose.append(fila_total2)
        ws_desglose.cell(row=ws_desglose.max_row, column=1).font = Font(bold=True)

        for col, ancho in zip(
            "ABCDEFGHIJKLMNOPQRST",
            (14, 12, 6, 28, 15, 10, 14, 14, 30, 24, 14, 20, 12, 14, 14, 6, 16, 14, 14, 12),
        ):
            ws_desglose.column_dimensions[col].width = ancho

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="reporte_vega.xlsx"'
        wb.save(response)
        return response

    return render(request, "core/reportes.html", context)


@login_required
def aplicar_utm(request):
    if request.user.perfil.rol != "caja":
        messages.error(request, "No tienes permisos para aplicar valor UTM.")
        return redirect("pagos")

    if request.method == "POST":
        try:
            valor_utm = Decimal(request.POST.get("valor_utm", "").strip())
            mes = int(request.POST.get("mes", "").strip())
            if valor_utm <= 0:
                raise ValueError("El valor UTM debe ser mayor que cero.")
            if not (1 <= mes <= 12):
                raise ValueError("Mes inválido, debe ser entre 1 y 12.")
            anio = date.today().year
            qs = Pago.objects.filter(fecha__year=anio, fecha__month=mes).select_related("local")
            contador = 0
            for pago in qs:
                pago.valor = pago.local.cantidad_utm * valor_utm
                pago.save()
                contador += 1
            ValorUTM.objects.update_or_create(
                anio=anio,
                mes=mes,
                defaults={"valor": valor_utm},
            )
            messages.success(
                request,
                f"Se actualizaron {contador} arriendos del mes {mes}/{anio} con valor UTM ${valor_utm:,.2f}.",
            )
        except (ValueError, InvalidOperation) as e:
            messages.error(request, f"No se pudo aplicar: {e}")
        return redirect("aplicar_utm")

    context = {
        "utm_base": UTM_BASE,
        "mes_actual": date.today().month,
        "anio_actual": date.today().year,
        "utm_por_mes": {
            f"{v.anio}-{v.mes}": v.valor
            for v in ValorUTM.objects.all()
        },
    }
    return render(request, "core/aplicar_utm.html", context)


@login_required
def guardar_utm(request):
    if request.user.perfil.rol != "caja":
        messages.error(request, "No tienes permisos para modificar el valor UTM.")
        return redirect("pagos")

    tipo = request.POST.get("tipo", "local")
    if request.method == "POST":
        try:
            mes = int(request.POST.get("mes", "").strip())
            valor = Decimal(request.POST.get("valor", "").strip())
            if valor <= 0:
                raise ValueError("El valor UTM debe ser mayor que cero.")
            if not (1 <= mes <= 12):
                raise ValueError("Mes inválido, debe ser entre 1 y 12.")
            anio = date.today().year
            ValorUTM.objects.update_or_create(
                anio=anio,
                mes=mes,
                defaults={"valor": valor},
            )
            messages.success(request, f"Valor UTM del mes {mes}/{anio} actualizado a ${valor:,.2f}.")
        except (ValueError, InvalidOperation) as e:
            messages.error(request, f"No se pudo guardar: {e}")
    return redirect("pagos_tipo", tipo=tipo)
