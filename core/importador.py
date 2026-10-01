import unicodedata
from datetime import date, datetime

from django.db import transaction

from core.models import Local, Pago


MESES_ES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]


def norm(s):
    if s is None:
        return ""
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().strip().lower()


def nombre_columa_limpio(s):
    n = norm(s)
    n = n.replace("n°", "numero").replace("no.", "numero").replace("#", "numero")
    n = n.replace("nro", "numero").replace("num", "numero")
    return n


def parsear_fecha(valor):
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s).date()
    except ValueError:
        return None


def parsear_decimal(valor):
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, (int, float)):
        return valor
    s = str(valor).replace("$", "").replace(".", "").replace(",", ".").strip()
    try:
        return float(s)
    except ValueError:
        return None


def parsear_mes(mes):
    if not mes:
        return ""
    n = norm(mes)
    for i, m in enumerate(MESES_ES, start=1):
        if norm(m) == n or norm(m)[:3] == n[:3]:
            return m
    return mes


def importar_workbook(ws, tipo_default=None):
    """Lee una hoja de Excel y crea/actualiza Pagos. Devuelve dict de contador."""
    contador = {"creados": 0, "actualizados": 0, "sin_cambios": 0, "errores": 0, "fila": []}

    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        contador["error"] = "El Excel está vacío."
        return contador

    headers = [nombre_columa_limpio(c) for c in filas[0]]
    idx = {
        "numero": next((i for i, h in enumerate(headers) if "numero" in h and "local" not in h and "folio" not in h), None),
        "local": next((i for i, h in enumerate(headers) if "local" in h and "numero" not in h), None),
        "empresario": next((i for i, h in enumerate(headers) if "empresario" in h or "nombre" in h), None),
        "rut": next((i for i, h in enumerate(headers) if "rut" in h), None),
        "descripcion": next((i for i, h in enumerate(headers) if "descripcion" in h), None),
        "tipo_deuda": next((i for i, h in enumerate(headers) if "deuda" in h or "tipo" in h), None),
        "domicilio": next((i for i, h in enumerate(headers) if "domicilio" in h or "direccion" in h), None),
        "fono": next((i for i, h in enumerate(headers) if "fono" in h or "telefono" in h or "telef" in h), None),
        "correo": next((i for i, h in enumerate(headers) if "correo" in h or "email" in h), None),
        "periodo": next((i for i, h in enumerate(headers) if "periodo" in h), None),
        "valor": next((i for i, h in enumerate(headers) if "valor" in h or "monto" in h), None),
        "fecha_pago": next((i for i, h in enumerate(headers) if "pago" in h and "fecha" in h), None),
        "fecha": next((i for i, h in enumerate(headers) if h == "fecha"), None),
        "mes": next((i for i, h in enumerate(headers) if h == "mes"), None),
        "anio": next((i for i, h in enumerate(headers) if h in ("anio", "año") or "anio" in h), None),
        "folio": next((i for i, h in enumerate(headers) if "folio" in h), None),
        "tipo": next(
            (i for i, h in enumerate(headers) if h == "tipo"),
            next((i for i, h in enumerate(headers) if "tipo" in h and "deuda" not in h), None),
        ),
    }
    if idx["numero"] is None and idx["local"] is None:
        contador["error"] = "No se encontró la columna 'N° Local'. Cabeceras: %s" % [h or "?" for h in headers]
        return contador

    for num_fila, fila in enumerate(filas[1:], start=2):
        try:
            numero_celda = fila[idx["numero"] if idx["numero"] is not None else idx["local"]]
            numero = str(numero_celda).strip() if numero_celda is not None else ""
            if not numero:
                if not any(c not in (None, "") for c in fila):
                    continue
                contador["errores"] += 1
                contador["fila"].append(f"Fila {num_fila}: N° Local vacío, se omitió.")
                continue

            tipo = fila[idx["tipo"]] if idx["tipo"] is not None else None
            tipo = (str(tipo).strip() if tipo else tipo_default) or "local"
            local, _ = Local.objects.get_or_create(
                numero=numero, defaults={"tipo": tipo, "ocupado": True}
            )
            if tipo and local.tipo != tipo:
                local.tipo = tipo
                local.save()

            fecha_pago_obj = parsear_fecha(fila[idx["fecha_pago"]] if idx["fecha_pago"] is not None else None)
            fecha_obj = parsear_fecha(fila[idx["fecha"]] if idx["fecha"] is not None else None)
            if fecha_obj is None and fecha_pago_obj is not None:
                fecha_obj = fecha_pago_obj
            if fecha_obj is None:
                anio = str(fila[idx["anio"]]).strip() if idx["anio"] is not None and fila[idx["anio"]] else None
                if anio and anio.isdigit():
                    fecha_obj = date(int(anio), 1, 1)

            mes_salida = parsear_mes(fila[idx["mes"]] if idx["mes"] is not None else None)
            if not mes_salida and fecha_obj:
                mes_salida = MESES_ES[fecha_obj.month - 1]
            periodo = str(fila[idx["periodo"]]).strip() if idx["periodo"] is not None and fila[idx["periodo"]] else ""
            if not periodo and fecha_obj:
                periodo = f"{mes_salida} {fecha_obj.year}"

            valor = parsear_decimal(fila[idx["valor"]] if idx["valor"] is not None else None)
            if valor is None:
                contador["errores"] += 1
                contador["fila"].append(f"Fila {num_fila}: valor inválido o faltante en local {numero}.")
                continue

            folio = str(fila[idx["folio"]]).strip() if idx["folio"] is not None and fila[idx["folio"]] else ""
            if not folio and fecha_obj:
                folio = f"{numero}-{fecha_obj.year}-{fecha_obj.month}"

            datos = {
                "fecha": fecha_obj,
                "mes": mes_salida,
                "empresario": str(fila[idx["empresario"]]).strip() if idx["empresario"] is not None and fila[idx["empresario"]] else "",
                "rut": str(fila[idx["rut"]]).strip() if idx["rut"] is not None and fila[idx["rut"]] else "",
                "descripcion": str(fila[idx["descripcion"]]).strip() if idx["descripcion"] is not None and fila[idx["descripcion"]] else "",
                "tipo_deuda": str(fila[idx["tipo_deuda"]]).strip() if idx["tipo_deuda"] is not None and fila[idx["tipo_deuda"]] else "",
                "domicilio": str(fila[idx["domicilio"]]).strip() if idx["domicilio"] is not None and fila[idx["domicilio"]] else "",
                "fono": str(fila[idx["fono"]]).strip() if idx["fono"] is not None and fila[idx["fono"]] else "",
                "correo": str(fila[idx["correo"]]).strip() if idx["correo"] is not None and fila[idx["correo"]] else "",
                "periodo": periodo,
                "valor": valor,
                "fecha_pago": fecha_pago_obj,
            }

            pago = Pago.objects.filter(folio=folio).first() if folio else None
            if pago is None:
                Pago.objects.create(local=local, folio=folio, **datos)
                contador["creados"] += 1
            else:
                cambiado = False
                for campo, valor_campo in datos.items():
                    if getattr(pago, campo) != valor_campo:
                        setattr(pago, campo, valor_campo)
                        cambiado = True
                if pago.local != local:
                    pago.local = local
                    cambiado = True
                if cambiado:
                    pago.save()
                    contador["actualizados"] += 1
                else:
                    contador["sin_cambios"] += 1

        except Exception as exc:
            contador["errores"] += 1
            contador["fila"].append(f"Fila {num_fila}: error {exc}")

    return contador


def importar_excel_desde_bytes(contenido, tipo_default=None):
    """Recibe bytes de un .xlsx y devuelve el dict de contador (o mensaje de error)."""
    from django.core.files.uploadedfile import InMemoryUploadedFile

    if isinstance(contenido, InMemoryUploadedFile):
        contenido = contenido.read()

    from io import BytesIO
    from openpyxl import load_workbook

    try:
        wb = load_workbook(BytesIO(contenido), data_only=True)
    except Exception as exc:
        return {"error": f"No se pudo leer el archivo como Excel: {exc}"}

    ws = wb.worksheets[0]
    return importar_workbook(ws, tipo_default)